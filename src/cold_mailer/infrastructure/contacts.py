"""Bounded CSV/XLSX import with row diagnostics and stable string values."""

from __future__ import annotations

import csv
import hashlib
import io
import re
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path

from cold_mailer.domain.models import ImportIssue, ImportReport, ValidationError, address

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10_000


@dataclass(frozen=True)
class ContactDocument:
    headers: tuple[str, ...]
    rows: tuple[tuple[object, ...], ...] = field(repr=False)
    sha256: str

    def parse(self, mapping: dict[str, str], suppressed: set[str] | None = None) -> ImportReport:
        columns = mapped_columns(self.headers, mapping)
        selected = [index for index, column in enumerate(columns) if column]
        # Retain malformed row widths so the existing row diagnostics still reject them.
        rows: list[list[object]] = [[columns[index] for index in selected]]
        for row in self.rows:
            rows.append(
                [row[index] for index in selected]
                if len(row) == len(columns)
                else (
                    ["invalid column count"] * (len(selected) + 1)
                    if any(str(cell or "").strip() for cell in row)
                    else []
                )
            )
        return parse_rows(rows, suppressed)


def mapped_columns(headers: tuple[str, ...], mapping: dict[str, str]) -> list[str]:
    if set(mapping) != set(headers):
        raise ValidationError("The mapping must match this file's exact header names.")
    columns = [mapping[header].strip().lower() for header in headers]
    active = [column for column in columns if column]
    if "email" not in active or len(active) != len(set(active)):
        raise ValidationError(
            "Map exactly one column to email and avoid duplicate destination fields."
        )
    if any(not re.fullmatch(r"[a-z][a-z0-9_]{0,63}", column) for column in active):
        raise ValidationError(
            "Destination fields must use letters, numbers and underscores, starting with a letter (64 characters maximum)."
        )
    return columns


def read_document(path: Path, *, expected_sha256: str | None = None) -> ContactDocument:
    if not path.is_file() or path.stat().st_size > MAX_BYTES:
        raise ValidationError("Choose a CSV/XLSX file no larger than 10 MiB.")
    with path.open("rb") as stream:
        raw = stream.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise ValidationError("The contact file grew beyond the 10 MiB limit.")
    checksum = hashlib.sha256(raw).hexdigest()
    if expected_sha256 is not None and checksum != expected_sha256:
        raise ValidationError(
            "The source file changed after preview. Preview it again before importing."
        )
    if path.suffix.lower() == ".csv":
        try:
            rows = []
            for row in csv.reader(io.StringIO(raw.decode("utf-8-sig"), newline="")):
                rows.append(row)
                if len(rows) > MAX_ROWS + 1:
                    raise ValidationError("At most 10,000 contact rows are supported.")
        except UnicodeError, csv.Error:
            raise ValidationError("CSV must be valid UTF-8 with a header row.") from None
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValidationError("The expanded workbook exceeds the 50 MiB limit.")
            workbook = load_workbook(
                io.BytesIO(raw), read_only=True, data_only=False, keep_links=False
            )
            try:
                if workbook.active is None:
                    raise ValidationError("The workbook contains no active sheet.")
                rows = []
                for row in workbook.active.values:
                    rows.append(list(row))
                    if len(rows) > MAX_ROWS + 1:
                        raise ValidationError("At most 10,000 contact rows are supported.")
            finally:
                workbook.close()
        except zipfile.BadZipFile, KeyError, ValueError:
            raise ValidationError("The workbook is invalid or unsupported.") from None
    else:
        raise ValidationError("Only CSV and XLSX are supported.")
    if not rows or len(rows) > MAX_ROWS + 1:
        raise ValidationError("A header and at most 10,000 contact rows are required.")
    headers = tuple(str(value or "").strip() for value in rows[0])
    if (
        not headers
        or len(headers) > 100
        or any(not header or len(header) > 200 for header in headers)
        or len({header.lower() for header in headers}) != len(headers)
    ):
        raise ValidationError("Use 1-100 unique, nonempty source headers (200 characters maximum).")
    return ContactDocument(headers, tuple(tuple(row) for row in rows[1:]), checksum)


def load_contacts(path: Path, suppressed: set[str] | None = None) -> ImportReport:
    if not path.is_file() or path.stat().st_size > MAX_BYTES:
        raise ValidationError("Choose a CSV/XLSX file no larger than 10 MiB.")
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            try:
                return parse_rows(csv.reader(stream), suppressed)
            except UnicodeError, csv.Error:
                raise ValidationError("CSV must be valid UTF-8 with a header row.") from None
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        try:
            with zipfile.ZipFile(path) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * MAX_BYTES:
                    raise ValidationError("The expanded workbook exceeds the size limit.")
            workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.active
                if sheet is None:
                    raise ValidationError("The workbook contains no active sheet.")
                return parse_rows(sheet.values, suppressed)
            finally:
                workbook.close()
        except zipfile.BadZipFile, KeyError, ValueError:
            raise ValidationError("The workbook is invalid or unsupported.") from None
    raise ValidationError("Only CSV and XLSX are supported; convert legacy XLS files first.")


def parse_rows(
    rows: Iterable[Iterable[object]], suppressed: set[str] | None = None
) -> ImportReport:
    iterator = iter(rows)
    columns = [str(cell or "").strip().lower() for cell in next(iterator, [])]
    if not columns or "email" not in columns or len(columns) != len(set(columns)) or "" in columns:
        raise ValidationError("A unique, nonempty header row containing email is required.")
    if len(columns) > 100:
        raise ValidationError("At most 100 contact columns are supported.")
    report = ImportReport()
    seen: set[str] = set()
    for number, row in enumerate(iterator, 2):
        if number > MAX_ROWS + 1:
            raise ValidationError("At most 10,000 contact rows are supported.")
        values = [
            "" if cell is None else str(cell).replace("\r\n", "\n").replace("\r", "\n").strip()
            for cell in row
        ]
        if not any(values):
            continue
        if len(values) != len(columns) or any(len(cell) > 10_000 for cell in values):
            report.issues.append(ImportIssue(number, "Column count or field length is invalid."))
            continue
        contact = dict(zip(columns, values, strict=True))
        try:
            contact["email"] = address(contact["email"])
        except ValidationError as exc:
            report.issues.append(ImportIssue(number, str(exc)))
            continue
        key = contact["email"]
        if key in (suppressed or set()):
            report.suppressed += 1
        elif key in seen:
            report.duplicates += 1
        else:
            report.contacts.append(contact)
            seen.add(key)
    return report


class DataLoader:
    """Compatibility entry point for callers of the original loader."""

    @staticmethod
    def load_contacts(file_path: str) -> list[dict[str, str]]:
        return load_contacts(Path(file_path)).contacts
